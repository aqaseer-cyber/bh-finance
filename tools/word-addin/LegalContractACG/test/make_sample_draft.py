#!/usr/bin/env python3
"""
Builds test/sample-draft.docx: the agreement text as a messy draft with
typed numbers ("1.", "1.1", "a)"), Calibri 11, no styles, no footer.
This is the input the add-in has to turn into the reference layout.

    python3 make_sample_draft.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "sample-draft.docx"

PARAS = [
    "CONSULTANCY AGREEMENT",
    "(Trial Engagement)",
    'This Consultancy Agreement (the "Agreement") is made on the date of the last signature below (the "Effective Date") between:',
    '(1) Almoayyed Contracting Group W.L.L., Commercial Registration No. 43090-1, of P.O. Box 32232, Manama, Kingdom of Bahrain (the "Client"); and',
    '(2) S2 International Consultants, Commercial Registration No. 194826-1, of Suite 106, Building 143, Block 340, Road 41, Al Juffair, Manama, Kingdom of Bahrain, represented by Mr. Jose (Joseph) Santos Esparza III, Chief Executive Officer (the "Consultant"; Mr. Esparza being the "Key Person").',
    'The Client and the Consultant are each a "Party" and together the "Parties".',
    "1. ENGAGEMENT AND SERVICES",
    '1.1 Engagement. The Client engages the Consultant, for the Trial Period defined in Clause 3.1, to provide the services in Clause 1.2 (the "Services") on the terms of this Agreement. The Trial Target as set out in Clause 1.4 is the agreed measure of the trial.',
    "1.2 Services. The Consultant shall:",
    "a) identify and introduce to the Client business opportunities with U.S. Government and military entities, prime and defense contractors in the Kingdom of Bahrain and the region, through the Opportunity Notice procedure in Clause 1.3;",
    "b) support the Client's tender and estimation department in preparing bids, proposals, capability statements and sources-sought responses, within the deadlines set by the Client;",
    "c) advise the Client's designated personnel, on request, on the Client's registrations on U.S. Government and defense procurement platforms (such as SAM, PIEE, JCCS and prime-contractor supplier portals); and",
    "d) facilitate introductions, meetings and communications with prospective customers and partners as directed or authorised by the Client, in compliance with Clause 1.6.",
    "1.3 Registered Opportunities.",
    '(a) The Consultant shall notify the Client in writing of each opportunity it proposes to pursue (an "Opportunity Notice").',
    '(b) Within ten (10) business days the Client shall confirm in writing whether it accepts the opportunity as a "Registered Opportunity".',
    "(c) No commission is payable on any opportunity that is not a Registered Opportunity.",
    '1.4 Trial Target. The Parties have set, as the measure of the Trial Period, binding awards to the Client of Registered Opportunities with an aggregate Net Contract Value of not less than BHD 300,000 (three hundred thousand Bahraini Dinars) received on or before the last day of the Trial Period (the "Trial Target").',
    "1.5 Reporting. The Consultant shall deliver to the Client's designated representative every two (2) weeks an updated register of all Opportunity Notices, and shall attend a review meeting at least monthly.",
    "1.6 Conduct of the Services. The Consultant shall:",
    "a) perform the Services personally through the Key Person and not subcontract or delegate them;",
    "b) not make any bid, proposal, representation or commitment in the Client's name except with the prior written approval of the Client's authorised representative; and",
    "c) conduct the Services lawfully, ethically and with integrity, in compliance with all applicable laws.",
    "1.7 Exclusivity. The Consultant's engagement under this Agreement is exclusive.",
    ("quote", "In the event that this contract terminates before its term or in the event that it is not renewed, the Consultant shall be obligated not to compete with the Client in any way whatsoever within the scope of the business in which the consultations were provided, and this restriction is valid for a complete year."),
    "2. FEES",
    "2.1 Retainer. The Client shall pay the Consultant a retainer of USD 2,500 (two thousand five hundred United States Dollars) per month of the Trial Period, monthly in arrears within thirty (30) days of a valid invoice.",
    "2.2 Commission. For each Registered Opportunity awarded to the Client during the Trial Period, the Client shall pay the Consultant a commission as set out in this Clause.",
    "2.3 Payment of commission.",
    "a) Commission is earned when the Client receives a binding award.",
    "b) Commission is payable back-to-back with the Client's receipt of payment from the customer.",
    "c) The Consultant shall keep records supporting each commission claim.",
    "2.4 Expenses. No expenses are reimbursable unless approved in writing in advance.",
    "3. TERM AND TERMINATION",
    "3.1 Trial Period. This Agreement runs for ninety (90) days from the Effective Date.",
    "3.2 Termination. Either Party may terminate this Agreement on fourteen (14) days' written notice.",
    "4. GENERAL",
    "4.1 Law and disputes. This Agreement is governed by the laws of the Kingdom of Bahrain.",
    "4.2 Miscellaneous. This Agreement, with its Schedule, is the entire agreement between the Parties.",
    "SIGNATURES",
    "Signed by the Parties' authorised representatives on the dates below.",
    "For the Client",
    "Name: Hala F. Almoayyed",
    "Title: Executive Director",
    "Signature: ______________________________",
    "Date: ______________________________",
    "Company stamp:",
    "For the Consultant",
    "Name: Mr. Jose (Joseph) Santos Esparza III",
    "Title: Chief Executive Officer",
    "Signature: ______________________________",
    "Date: ______________________________",
    "Company stamp:",
    "SCHEDULE 1",
    "Client's pre-existing customers, agencies, prime contractors and projects (Clause 1.3(c)). No commission is payable in respect of any of them.",
    ("table", [
        ["No.", "Entity / agency / prime contractor", "Nature of pre-existing relationship"],
        ["1", "KBR Services (Kellogg Brown and Root Services)", "Subcontract"],
        ["2", "AiLog International", "Joint Venture"],
        ["3", "Eco Flow Purification W.L.L", "Subcontract"],
    ]),
]


def main() -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for item in PARAS:
        if isinstance(item, tuple) and item[0] == "table":
            rows = item[1]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, txt in enumerate(row):
                    t.cell(r, c).text = txt
            continue
        if isinstance(item, tuple) and item[0] == "quote":
            p = doc.add_paragraph(item[1])
            p.paragraph_format.left_indent = Pt(36)
            continue
        doc.add_paragraph(item)
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
