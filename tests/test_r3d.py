"""v3 R3d — the consistency contract: unified artifact naming, the
shell's run stamp, export_run's four-file bundle + manifest, and the
cross-artifact FV/MoS abort (a test corrupts one artifact's FV and
expects the abort — the design gate). All offline.
"""
import datetime as dt
import hashlib
import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import build_testco_companyfacts
from forensic_viz.edgar import parse_companyfacts
from forensic_viz.export import (
    ConsistencyError, assert_run_consistency, export_run,
)
from forensic_viz.metrics import (
    DashboardData, apply_track, build_fundamental_metrics,
    build_price_metrics,
)
from forensic_viz.prices import PriceSeries
from forensic_viz.runid import artifact_name, artifact_stem, run_identity
from forensic_viz.valuation import CaseInputs, ValuationInputs, build_valuation
from forensic_viz.verdict import build_verdict

FIXTURES = Path(__file__).parent / "fixtures"

STEM_RX = re.compile(r"^TESTCO_\d{4}-\d{2}-\d{2}_[0-9a-f]{8}$")


def _testco(prices=True):
    d = DashboardData(ticker="TESTCO", company="TESTCO INC", subtitle="fx",
                      generated=dt.date(2026, 7, 19))
    d.sic_code = "3571"
    apply_track(d, "auto")
    build_fundamental_metrics(
        parse_companyfacts(build_testco_companyfacts(), "TESTCO"), d)
    if prices:
        raw = json.loads((FIXTURES / "aapl_weekly_5y.json").read_text())
        build_price_metrics(PriceSeries(
            symbol="TESTCO",
            dates=[dt.date.fromisoformat(s) for s in raw["dates"]],
            closes=raw["close"], source="fixture"), d)
    d.thesis, d.terminal_risk = "t", "r"
    return d


def _valued(d):
    inputs = ValuationInputs(
        method="dcf", discount_rate=0.09,
        cases={"Bear": CaseInputs(g0=0.02, g_term=0.02),
               "Base": CaseInputs(g0=0.05, g_term=0.025),
               "Bull": CaseInputs(g0=0.08, g_term=0.03)})
    res = build_valuation(d, inputs)
    return res, build_verdict(d, inputs, res, rating="Buy")


# ------------------------------------------------------------- naming

def test_artifact_naming_contract():
    d = _testco(prices=False)
    stem = artifact_stem(d)
    assert STEM_RX.match(stem)
    rid, _ = run_identity(d)
    assert stem.endswith(rid)
    assert artifact_name(d, None, "report") == f"{stem}_report.pdf"
    assert artifact_name(d, None, "model") == f"{stem}_model.xlsx"
    assert artifact_name(d, None, "shell") == f"{stem}_shell.xlsx"
    assert artifact_name(d, None, "manifest") == f"{stem}_manifest.json"
    with pytest.raises(ValueError):
        artifact_name(d, None, "csv")


# ------------------------------------------------- the four-file bundle

def test_export_run_produces_exactly_four_files(tmp_path):
    d = _testco()
    res, v = _valued(d)
    manifest = export_run(d, res=res, verdict=v, out_dir=str(tmp_path))
    files = sorted(p.name for p in tmp_path.iterdir())
    assert len(files) == 4                       # the design gate
    stems = {f.rsplit("_", 1)[0] for f in files}
    assert len(stems) == 1 and STEM_RX.match(stems.pop())
    # the manifest is the machine-readable receipt
    on_disk = json.loads(
        Path(manifest["manifest_path"]).read_text(encoding="utf-8"))
    assert on_disk["run_id"] == manifest["run_id"] == run_identity(d, res)[0]
    assert on_disk["rating"] == "Buy"
    assert on_disk["fv_avg"] == pytest.approx(v.fv_avg)
    assert on_disk["providers"].startswith("EDGAR")
    assert on_disk["warnings"] >= len(d.health_notes or [])
    for kind in ("report", "model", "shell"):
        art = on_disk["artifacts"][kind]
        p = Path(art["path"])
        assert p.exists()
        assert hashlib.sha256(p.read_bytes()).hexdigest() == art["sha256"]


def test_export_run_without_valuation_declares_the_missing_shell(tmp_path):
    d = _testco(prices=False)
    manifest = export_run(d, out_dir=str(tmp_path))
    assert len(list(tmp_path.iterdir())) == 3    # report, model, manifest
    assert "shell" not in manifest["artifacts"]
    assert "no valuation attached" in manifest["shell"]
    assert manifest["fv_avg"] is None and manifest["rating"] is None


# ------------------------------------------------- the shell run stamp

def test_shell_stamp_carries_run_identity_and_fv(tmp_path):
    from forensic_viz.workbook import fill_workbook
    d = _testco()
    res, v = _valued(d)
    out = tmp_path / "shell.xlsx"
    fill_workbook(d, str(out), res=res, verdict=v)
    cell = load_workbook(str(out))["Control"]["B7"]
    assert cell.comment is not None
    rid, ihash = run_identity(d, res)
    text = cell.comment.text
    assert f"Run {rid}" in text and f"inputs {ihash}" in text
    assert "generated 2026-07-19" in text and "app " in text
    assert f"FV_avg {v.fv_avg!r}" in text


# --------------------------------------------- the consistency contract

def test_corrupted_model_fv_aborts_with_the_diff(tmp_path):
    d = _testco()
    res, v = _valued(d)
    manifest = export_run(d, res=res, verdict=v, out_dir=str(tmp_path))
    model = Path(manifest["artifacts"]["model"]["path"])
    shell = Path(manifest["artifacts"]["shell"]["path"])
    assert_run_consistency(v, model, shell)      # clean run passes
    wb = load_workbook(str(model))
    cov = wb["Cover"]
    for r in range(1, cov.max_row + 1):
        if cov.cell(row=r, column=1).value == "FV average":
            cov.cell(row=r, column=2).value = 999.99
    wb.save(str(model))
    with pytest.raises(ConsistencyError) as exc:
        assert_run_consistency(v, model, shell)
    assert "model Cover FV average: 999.99" in str(exc.value)
    assert "disagree on the verdict numbers" in str(exc.value)


def test_corrupted_shell_stamp_aborts_too(tmp_path):
    from openpyxl.comments import Comment
    d = _testco()
    res, v = _valued(d)
    manifest = export_run(d, res=res, verdict=v, out_dir=str(tmp_path))
    shell = Path(manifest["artifacts"]["shell"]["path"])
    wb = load_workbook(str(shell))
    wb["Control"]["B7"].comment = Comment(
        "Run deadbeef · inputs 0000000000 · FV_avg 1.0 · MoS 0.5",
        "corruptor")
    wb.save(str(shell))
    with pytest.raises(ConsistencyError):
        assert_run_consistency(v, shell_path=shell)


# -------------------------------------------------- the web run bundle

def test_api_export_run_kind(tmp_path):
    from fastapi.testclient import TestClient
    from test_api_contract import TOKEN, _sse_events, fixture_pipeline
    from webui.server import create_app

    app = create_app(pipeline=fixture_pipeline, token=TOKEN)
    client = TestClient(app, headers={"Authorization": f"Bearer {TOKEN}"})
    _sse_events(client, "/api/run/TESTCO")
    r = client.post("/api/export/run",
                    json={"ticker": "TESTCO", "out_dir": str(tmp_path)})
    assert r.status_code == 200
    data = r.json()["data"]
    manifest = data["manifest"]
    assert Path(data["path"]).exists()
    # no valuation posted in this session -> three files, declared
    assert "no valuation attached" in manifest["shell"]
    assert len(list(tmp_path.iterdir())) == 3
